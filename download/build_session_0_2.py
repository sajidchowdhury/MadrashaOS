"""
MadrashaOS - Session 0.2 Deliverable: Personas & Role Journeys
Output: .docx via python-docx (landscape for journey maps)
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = "/home/z/my-project/download/MadrashaOS_Session_0.2_Personas_Journeys_TechnicalDoc_2026-09-16.docx"

# ---------- Helpers ----------

def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'BFBFBF')
        borders.append(e)
    tbl_pr.append(borders)

def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for name, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        m = OxmlElement(f'w:{name}'); m.set(qn('w:w'), str(val)); m.set(qn('w:type'), 'dxa'); tc_mar.append(m)
    tc_pr.append(tc_mar)

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text); run.bold = True
    if level == 1:
        run.font.size = Pt(16); run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    elif level == 2:
        run.font.size = Pt(13); run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    else:
        run.font.size = Pt(11.5); run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def add_para(doc, text, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text); run.font.size = Pt(size); run.italic = italic
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    if p.runs: p.runs[0].text = ''
    r = p.add_run(text); r.font.size = Pt(11)
    return p

def build_table(doc, headers, rows, col_widths_cm=None, header_bg='1F3A5F', zebra=True, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    set_table_borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_margins(c); set_cell_bg(c, header_bg)
        c.text = ''
        p = c.paragraphs[0]
        r = p.add_run(h); r.bold = True; r.font.size = Pt(font_size + 0.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        if col_widths_cm: c.width = Cm(col_widths_cm[i])
    for ri, row in enumerate(rows, start=1):
        new_row = t.add_row()
        for ci, val in enumerate(row):
            c = new_row.cells[ci]
            set_cell_margins(c)
            if zebra and ri % 2 == 0: set_cell_bg(c, 'F7F7F7')
            c.text = ''
            p = c.paragraphs[0]
            r = p.add_run(str(val)); r.font.size = Pt(font_size)
            if col_widths_cm: c.width = Cm(col_widths_cm[ci])
    return t

def add_persona_card(doc, persona):
    add_heading(doc, f"{persona['no']} — {persona['role']}", 3)
    fields = [
        ('Name (placeholder)', persona['name']),
        ('Demographic', persona['demographic']),
        ('Tech Literacy', persona['tech_literacy']),
        ('Top-3 Daily Tasks', persona['top_tasks']),
        ('Goals', persona['goals']),
        ('Frustrations', persona['frustrations']),
        ('Key Permissions (SRS)', persona['permissions']),
        ('Mobile Usage', persona['mobile_use']),
        ('SRS Ref', persona['srs_ref']),
        ('Quote (placeholder)', persona['quote']),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    set_table_borders(t)
    for k, v in fields:
        row = t.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_margins(c0); set_cell_margins(c1)
        set_cell_bg(c0, 'F2F2F2')
        c0.text = ''; c1.text = ''
        r0 = c0.paragraphs[0].add_run(k); r0.bold = True; r0.font.size = Pt(9.5)
        r1 = c1.paragraphs[0].add_run(v); r1.font.size = Pt(9.5)
        c0.width = Cm(5); c1.width = Cm(21)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- Build Document ----------

doc = Document()
for section in doc.sections:
    section.page_width = Cm(29.7); section.page_height = Cm(21)
    section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2); section.right_margin = Cm(2)

style = doc.styles['Normal']
style.font.name = 'Calibri'; style.font.size = Pt(11)

# Title
title = doc.add_paragraph()
t_run = title.add_run('MadrashaOS — Session 0.2 Deliverable')
t_run.bold = True; t_run.font.size = Pt(20); t_run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(8)
s_run = sub.add_run('Personas & Role Journeys')
s_run.font.size = Pt(13); s_run.italic = True; s_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

# Meta table
meta = doc.add_table(rows=5, cols=2)
set_table_borders(meta)
meta_rows = [
    ('Document', 'MadrashaOS_Session_0.2_Personas_Journeys_TechnicalDoc_2026-09-16'),
    ('Source', 'SRS v2.0 (2026-09-16) \u00a78.1 role definitions; \u00a72 module specs'),
    ('Session', '0.2 \u2014 Personas & Role Journeys'),
    ('Phase', '0 \u2014 Discovery & Foundations'),
    ('Exit Criteria', 'Client signs off the 8 personas and 4 journey maps'),
]
for i, (k, v) in enumerate(meta_rows):
    c0 = meta.rows[i].cells[0]; c1 = meta.rows[i].cells[1]
    set_cell_margins(c0); set_cell_margins(c1); set_cell_bg(c0, 'F2F2F2')
    c0.text = ''; c1.text = ''
    r0 = c0.paragraphs[0].add_run(k); r0.bold = True; r0.font.size = Pt(10)
    r1 = c1.paragraphs[0].add_run(v); r1.font.size = Pt(10)
    c0.width = Cm(4); c1.width = Cm(22)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# 1. Overview
add_heading(doc, '1. Overview', 1)
add_para(doc,
    "This document scaffolds the 8 personas and 4 journey maps required by Phase 0 Session 0.2. "
    "Role definitions, permissions, and primary tasks are derived directly from SRS \u00a78.1 and the "
    "Part 2 module specs. Fields that require client input \u2014 real interviewee names, demographic "
    "detail, verbatim quotes, and confirmed pain points \u2014 are marked [placeholder]. These are filled "
    "during the 2-hour stakeholder interview block scheduled by the client. Once filled, the client "
    "signs off the personas and journey maps; that sign-off is the Session 0.2 exit criterion.")

# 2. Personas
add_heading(doc, '2. Personas', 1)
add_para(doc,
    "Eight personas cover every role defined in SRS \u00a78.1. Each persona card captures demographic, "
    "tech literacy, top-3 daily tasks (cross-referenced to the Session 0.1 module taxonomy), goals, "
    "frustrations, key permissions, and mobile usage pattern. Cards are designed to be printable on "
    "a single A4 sheet each for stakeholder review.")

personas = [
    {
        'no': '2.1', 'role': 'Super Admin (Platform Operator)',
        'name': '[Interviewee TBD \u2014 platform operator staff]',
        'demographic': 'Male/Female, 30\u201345, IT/systems background, manages the SaaS platform',
        'tech_literacy': 'High',
        'top_tasks': '1. Provision new tenants (SRS 2.7.2)\n2. Configure platform security policy (2.1.5)\n3. Monitor tenant health + quota (2.7.2)',
        'goals': 'Guarantee tenant isolation; maintain platform uptime; support tenant rollouts',
        'frustrations': 'Cross-tenant debugging is opaque; no single pane of glass for tenant health',
        'permissions': 'Platform-wide ALL; not bound to a single tenant',
        'mobile_use': 'Rarely \u2014 desktop-first admin console',
        'srs_ref': 'SRS \u00a78.1, \u00a72.7.2',
        'quote': '[Quote TBD] \u201cI need to know immediately when a tenant crosses quota.\u201d',
    },
    {
        'no': '2.2', 'role': 'Authority (Madrasha Head / Principal)',
        'name': '[Interviewee TBD \u2014 madrasha principal]',
        'demographic': 'Male, 45\u201360, religious + administrative leadership, decides financial approvals',
        'tech_literacy': 'Low\u2013Medium',
        'top_tasks': '1. Approve expenses + purchases (2.7.1)\n2. View authority dashboard (2.6.4)\n3. Sign off results publication (2.3.4)',
        'goals': 'Institutional accountability; financial control; student outcomes visibility',
        'frustrations': 'Pending approvals pile up; dashboard figures feel stale; hard to compare branches',
        'permissions': 'All-branch, all-module read; approve on routing rules',
        'mobile_use': 'Moderate \u2014 approves on tablet during travel',
        'srs_ref': 'SRS \u00a78.1, \u00a72.7.1, \u00a72.6.4',
        'quote': '[Quote TBD] \u201cI should not have to chase people for approvals.\u201d',
    },
    {
        'no': '2.3', 'role': 'Administrator',
        'name': '[Interviewee TBD \u2014 madrasha office administrator]',
        'demographic': 'Male/Female, 28\u201345, office management background',
        'tech_literacy': 'Medium',
        'top_tasks': '1. Manage students + admissions (2.2.1, 2.2.2)\n2. Configure organization + branches (2.1.1)\n3. Manage users + roles (2.1.3)',
        'goals': 'Operational efficiency; correct configuration; no broken references',
        'frustrations': 'Promotions destroy history if done wrong; role permission matrix is hard to scan',
        'permissions': 'Tenant-wide CRUD except financial posting + cross-tenant ops',
        'mobile_use': 'Low\u2013Moderate \u2014 mostly desktop',
        'srs_ref': 'SRS \u00a78.1, \u00a72.1, \u00a72.2',
        'quote': '[Quote TBD] \u201cIf I disable Hostel by mistake, every dependent breaks.\u201d',
    },
    {
        'no': '2.4', 'role': 'Accountant',
        'name': '[Interviewee TBD \u2014 madrasha accountant]',
        'demographic': 'Male, 35\u201355, accounting background, BDT/ledger fluent',
        'tech_literacy': 'Medium',
        'top_tasks': '1. Collect fees + issue receipts (2.4.1)\n2. Record expenses + income (2.4.3)\n3. Reconcile ledger + cash/bank (2.4.4, 2.4.3)',
        'goals': 'Books balanced; receipts numbered; Zakat fund isolated; audit clean',
        'frustrations': 'Reconciliation to cash/bank is manual; outstanding fees hard to chase by month',
        'permissions': 'Finance module CRUD; cannot edit student academic info',
        'mobile_use': 'Moderate \u2014 collects fees on phone at the counter',
        'srs_ref': 'SRS \u00a78.1, \u00a72.4, \u00a73.6',
        'quote': '[Quote TBD] \u201cEvery taka must post to the ledger the moment I receive it.\u201d',
    },
    {
        'no': '2.5', 'role': 'Teacher',
        'name': '[Interviewee TBD \u2014 full-time madrasha teacher]',
        'demographic': 'Male/Female, 25\u201350, subject teacher for one or more classes',
        'tech_literacy': 'Low\u2013Medium',
        'top_tasks': '1. Take attendance (mobile, <60s) (2.3.2)\n2. Enter marks (2.3.3)\n3. View own classes + assignments (2.2.5, 2.6.4)',
        'goals': 'Spend less time on paperwork; correct marks; quick attendance',
        'frustrations': 'Phone is slow; roster scroll for 40 students; network drops mid-submit',
        'permissions': 'Own classes + sections; cannot view financial data or other teachers\u2019 marks',
        'mobile_use': 'High \u2014 attendance + marks on phone daily',
        'srs_ref': 'SRS \u00a78.1, \u00a72.3.2, \u00a72.3.3',
        'quote': '[Quote TBD] \u201cIf attendance takes more than a minute, I stop using it.\u201d',
    },
    {
        'no': '2.6', 'role': 'Storekeeper',
        'name': '[Interviewee TBD \u2014 madrasha store/inventory staff]',
        'demographic': 'Male, 30\u201350, store/inventory background',
        'tech_literacy': 'Low\u2013Medium',
        'top_tasks': '1. Receive stock from purchases (2.5.1, 2.5.2)\n2. Issue stock to consumers (2.5.1)\n3. Track low-stock alerts (2.6.4)',
        'goals': 'No stockouts; correct movement history; clean supplier outstanding',
        'frustrations': 'Inventory count drift; paper-based issue slips; low-stock discovered too late',
        'permissions': 'Inventory + purchase CRUD; cannot view finance or student data',
        'mobile_use': 'Moderate \u2014 issues stock on phone at the store',
        'srs_ref': 'SRS \u00a78.1, \u00a72.5.1\u20132.5.3',
        'quote': '[Quote TBD] \u201cI want a low-stock alert before I run out, not after.\u201d',
    },
    {
        'no': '2.7', 'role': 'Guardian (Parent)',
        'name': '[Interviewee TBD \u2014 guardian of one or more students]',
        'demographic': 'Male/Female, 30\u201355, parent, working professional or homemaker',
        'tech_literacy': 'Medium',
        'top_tasks': '1. View child attendance + results (2.2.3)\n2. View + pay outstanding fees (2.2.3, 2.4.1)\n3. Read notices from madrasha (2.6.1)',
        'goals': 'Stay informed about child; pay fees without visiting office; trust in institution',
        'frustrations': 'Have to call office for receipts; cannot see results before parent-teacher meeting; multi-child switching is awkward',
        'permissions': 'Read-only on own linked children; no edit on academic or financial data',
        'mobile_use': 'High \u2014 portal is mobile-first',
        'srs_ref': 'SRS \u00a78.1, \u00a72.2.3, \u00a75.4',
        'quote': '[Quote TBD] \u201cI want to see my child\u2019s attendance without calling the office.\u201d',
    },
    {
        'no': '2.8', 'role': 'Student (Optional Role)',
        'name': '[Interviewee TBD \u2014 senior student representative]',
        'demographic': 'Male/Female, 14\u201322, senior madrasha student',
        'tech_literacy': 'Medium\u2013High (digital native)',
        'top_tasks': '1. View own attendance + results (read-only)\n2. View own fee plan (2.4.1)\n3. Read notices (2.6.1)',
        'goals': 'Track own progress; understand fee obligations; not miss notices',
        'frustrations': 'No self-service view; have to ask guardian or office',
        'permissions': 'Read-only on own records only; no peer visibility',
        'mobile_use': 'High \u2014 mobile-only access',
        'srs_ref': 'SRS \u00a78.1 (optional role), \u00a72.2.1',
        'quote': '[Quote TBD] \u201cI want to check my own result before my parents ask.\u201d',
    },
]

for p in personas:
    add_persona_card(doc, p)

# 3. Role Journeys
add_heading(doc, '3. Role Journey Maps', 1)
add_para(doc,
    "Four journeys stress-test the persona task sequences against real interaction flows. Each journey "
    "is staged: Trigger \u2192 Pre-action \u2192 Action \u2192 Confirmation \u2192 Post-action. Touchpoints reference the "
    "Session 0.1 module IDs. Pain points link back to the UX Risk Register (R1\u2013R16).")

journey_headers = ['Stage', 'What user does', 'UI screen', 'API endpoint', 'Emotion', 'Pain point', 'Opportunity']

# Journey 1: Take Attendance (Teacher)
add_heading(doc, '3.1 Journey 1 \u2014 Take Attendance (Teacher, mobile)', 2)
add_para(doc, "Scenario: 08:30 bell, Class 5 Section A, 40 students, on a 5-inch phone over mobile data.", italic=True)
j1 = [
    ('Trigger', 'Bell rings; teacher opens app from phone home', 'Login \u2192 Teacher Dashboard', 'POST /auth/login', 'Neutral \u2013 mild stress', 'Cold start >3s on slow network', 'Skeleton loading <1s; remember-me token'),
    ('Pre-action', 'Tap \u201cTake Attendance\u201d card on dashboard', 'Teacher Dashboard widget', 'GET /teacher/classes/active', 'Mild interest', 'Multiple taps to reach class', 'Default to last-used class+section'),
    ('Action', 'Mark roster: default Present; tap to cycle Absent/Late/Leave', 'Attendance quick-toggle screen', 'GET /classes/{id}/students', 'Focused', 'Scroll fatigue for 40 students', 'Sticky Present-all + exception swipe'),
    ('Submit', 'Tap Submit; confirm dialog', 'Confirm dialog', 'POST /attendance/sessions + records', 'Relief \u2192 anxiety on flaky net', 'Network drops mid-submit', 'Optimistic local state + Idempotency-Key (SRS \u00a76.5)'),
    ('Post', 'Toast \u201cSubmitted \u2014 38 Present, 2 Absent\u201d; return to dashboard', 'Toast + Dashboard', '\u2014', 'Satisfaction', 'No undo path', '30s undo window + audit flag (R6)'),
]
build_table(doc, journey_headers, j1, col_widths_cm=[2.2, 4.5, 3.5, 4.0, 2.5, 3.5, 4.0], font_size=8.5)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# Journey 2: Collect Fees (Accountant)
add_heading(doc, '3.2 Journey 2 \u2014 Collect Fee Payment (Accountant, mobile or desktop)', 2)
add_para(doc, "Scenario: Guardian pays ৳5000 tuition at the counter for March installment.", italic=True)
j2 = [
    ('Trigger', 'Guardian arrives with cash; accountant opens Fee module', 'Fee Module landing', 'GET /fees/outstanding', 'Neutral', 'Outstanding report slow on big roster', 'Cached outstanding report'),
    ('Pre-action', 'Search student by partial mobile or code', 'Student fee plan view', 'GET /students?q=01...', 'Mild interest', 'Search returns too many matches', 'Fuzzy match + recent-students chip'),
    ('Action', 'Select installment \u2192 amount \u2192 method (Cash) \u2192 account', 'Collect Payment form (3-step)', '\u2014', 'Focused', 'Method/account picker clunky', 'Default to last-used Cash account'),
    ('Submit', 'Confirm + Submit; receipt generated', 'Receipt preview', 'POST /fees/payments (Idempotency-Key)', 'Relief', 'Duplicate submit risk', 'Idempotency-Key enforced (SRS \u00a76.5)'),
    ('Post', 'Print receipt PDF; ledger posts balanced entry', 'Receipt PDF + Toast', 'POST /ledger/entries (Golden Flow \u00a73.6)', 'Satisfaction', 'Receipt not branded yet (R13)', 'Branded PDF template (Phase 5.1)'),
]
build_table(doc, journey_headers, j2, col_widths_cm=[2.2, 4.5, 3.5, 4.0, 2.5, 3.5, 4.0], font_size=8.5)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# Journey 3: Record Expense (Accountant)
add_heading(doc, '3.3 Journey 3 \u2014 Record Expense (Accountant)', 2)
add_para(doc, "Scenario: ৳25000 maintenance expense; crosses approval threshold.", italic=True)
j3 = [
    ('Trigger', 'Vendor invoice received; accountant opens Accounting', 'Accounting module landing', 'GET /ledger/recent', 'Neutral', 'Hard to find right cost center', 'Default cost center per user'),
    ('Pre-action', 'Pick expense account + cost center + date', 'Record Expense form', '\u2014', 'Mild interest', 'Cost center list long', 'Recent + starred accounts at top'),
    ('Action', 'Enter ৳25000 + narration + attachment', 'Expense form (cont.)', '\u2014', 'Focused', 'Threshold routing unclear to user', 'Inline \u201cPending approval required\u201d hint (R8)'),
    ('Submit', 'Tap Submit; entry enters pending state', 'Pending confirmation', 'POST /ledger/entries (status=pending)', 'Relief \u2192 mild confusion', 'User thinks it posted', 'Toast: \u201cSent to approver\u201d + link'),
    ('Post', 'Approval queue shows entry; account balance unchanged until approved', 'Approval Queue', 'GET /approvals?role=accountant', 'Wait', 'No visibility on approver SLA', 'Show approver name + ETA (R15)'),
]
build_table(doc, journey_headers, j3, col_widths_cm=[2.2, 4.5, 3.5, 4.0, 2.5, 3.5, 4.0], font_size=8.5)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# Journey 4: View Child Results (Guardian)
add_heading(doc, '3.4 Journey 4 \u2014 View Child Results (Guardian, mobile)', 2)
add_para(doc, "Scenario: Guardian logs in after result publication to view Class 5 child mark sheet.", italic=True)
j4 = [
    ('Trigger', 'SMS notice: results published; guardian opens portal', 'Guardian Portal login', 'POST /auth/login', 'Anticipation', 'Login MFA friction', 'Trusted-device MFA skip'),
    ('Pre-action', 'Tap Results tab on guardian portal', 'Guardian Portal home', 'GET /guardian/children/{id}/results', 'Anticipation', 'Pre-publication tab empty \u2014 confusing', 'Empty state: \u201cResults publish on [date]\u201d (R7)'),
    ('Action', 'View mark sheet: subjects, marks, grade, GPA', 'Student Result view', 'GET /students/{id}/results', 'Focused \u2192 happy/anxious', 'Ranking column off = no position', 'Conditional column (R7)'),
    ('Submit', '\u2014 (read-only, no submit)', '\u2014', '\u2014', '\u2014', '\u2014', '\u2014'),
    ('Post', 'Tap Print Mark Sheet; save or share PDF', 'Print PDF', 'POST /pdf/jobs (result sheet)', 'Satisfaction', 'PDF not branded (R13)', 'Branded PDF + share-via-WhatsApp intent'),
]
build_table(doc, journey_headers, j4, col_widths_cm=[2.2, 4.5, 3.5, 4.0, 2.5, 3.5, 4.0], font_size=8.5)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# 4. Exit criteria check
add_heading(doc, '4. Session 0.2 Exit Criteria Check', 1)
add_para(doc,
    "Per the UI/UX Implementation Plan, Session 0.2 is complete when the client signs off the 8 personas "
    "and 4 journey maps. The scaffolds above cover every role in SRS \u00a78.1 and reference the Session 0.1 "
    "module taxonomy for task traceability. Pending client interview block: 2 hours, one representative "
    "per role. Once interviews are transcribed, the [placeholder] fields \u2014 real names, verbatim quotes, "
    "demographic detail, and confirmed frustrations \u2014 are filled in and the document is re-issued for "
    "client sign-off. Journey pain points already cross-reference the UX Risk Register (R6, R7, R8, R13, "
    "R15), so design decisions in Phase 0 Session 0.4 will resolve them.")

# 5. Next session
add_heading(doc, '5. Next Session: 0.3 \u2014 Heuristic Review of Competing Madrasha ERPs', 1)
add_para(doc,
    "Session 0.3 benchmarks 3 existing madrasha/school ERPs against 10 usability heuristics to produce a "
    "\u201cDo-not-do\u201d list of patterns to avoid. Inputs: public demo links of 3 competing products. Required "
    "from client: list 2\u20133 ERPs they have previously used or evaluated, plus screenshots of any in-house "
    "spreadsheets currently in use (these are the strongest signal for what the new system must NOT do).")

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
