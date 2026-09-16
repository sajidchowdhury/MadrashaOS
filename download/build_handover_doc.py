"""
MadrashaOS Implementation Handover Sequence - Technical Documentation
Generates a .docx file using python-docx.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


OUTPUT = "/home/z/my-project/download/MadrashaOS_HandoverSequence_TechnicalDoc_2026-09-16.docx"

# ---------- Helpers ----------

def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '6')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'BFBFBF')
        borders.append(e)
    tbl_pr.append(borders)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for name, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        m = OxmlElement(f'w:{name}')
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        tc_mar.append(m)
    tc_pr.append(tc_mar)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    else:
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def add_para(doc, text, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    run = p.runs[0] if p.runs else p.add_run('')
    p.runs[0].text = ''
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


# ---------- Build Document ----------

doc = Document()

# Page setup: A4, margins per technical documentation spec
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

# Base font
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# ----- Title block -----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(2)
t_run = title.add_run('MadrashaOS — Implementation Handover Sequence')
t_run.bold = True
t_run.font.size = Pt(20)
t_run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(10)
s_run = sub.add_run('Technical Documentation — Role Handover Order for SRS v2.0')
s_run.font.size = Pt(12)
s_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
s_run.italic = True

# Meta table (document control)
meta = doc.add_table(rows=4, cols=2)
meta.autofit = True
set_table_borders(meta)
meta_rows = [
    ('Document', 'MadrashaOS_HandoverSequence_TechnicalDoc_2026-09-16'),
    ('Source SRS', 'MadrashaOS Implementation-Grade SRS v2.0 (2026-09-16)'),
    ('Audience', 'Project Manager, Tech Lead, UI/UX Lead, Backend Lead, Frontend Lead'),
    ('Status', 'Approved for handover planning'),
]
for i, (k, v) in enumerate(meta_rows):
    c0 = meta.rows[i].cells[0]
    c1 = meta.rows[i].cells[1]
    set_cell_margins(c0); set_cell_margins(c1)
    set_cell_bg(c0, 'F2F2F2')
    c0.text = ''
    p0 = c0.paragraphs[0]
    r0 = p0.add_run(k)
    r0.bold = True
    r0.font.size = Pt(10)
    c1.text = ''
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(v)
    r1.font.size = Pt(10)
    c0.width = Cm(4.5)
    c1.width = Cm(12)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ----- 1. Overview -----
add_heading(doc, '1. Overview', 1)
add_para(doc,
    "This document defines the recommended order in which the approved MadrashaOS SRS v2.0 is "
    "handed to the three implementation roles: UI/UX Designer, Backend Developer, and Frontend "
    "Developer. The sequence is derived directly from the structure, dependencies, and open items "
    "of the SRS itself — not from generic assumptions. Each role receives a clear handoff artifact "
    "from the preceding role, and the order minimizes rework, mock-data throwaway, and integration "
    "friction. The recommendation applies to the first build pass of Phase 0 (Foundation modules) "
    "and Phase 1; subsequent phases may parallelize once the foundation contracts are stable.")

# ----- 2. Recommended Sequence (summary table) -----
add_heading(doc, '2. Recommended Handover Sequence', 1)
add_para(doc,
    "The single recommended order is: UI/UX Designer first, Backend Developer second, Frontend "
    "Developer third. The summary below captures the rationale, estimated duration, and the "
    "artifacts each role must produce before the next role can start.")

seq = doc.add_table(rows=4, cols=4)
set_table_borders(seq)
hdr = ['Order', 'Role', 'Est. Duration', 'Primary Handoff Artifact']
for i, h in enumerate(hdr):
    c = seq.rows[0].cells[i]
    set_cell_margins(c)
    set_cell_bg(c, '1F3A5F')
    c.text = ''
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

seq_data = [
    ('1', 'UI/UX Designer', '2–4 weeks', 'Design system, wireframes, hi-fi mockups, PDF/print templates, brand kit'),
    ('2', 'Backend Developer', '4–6 weeks (Phase 0)', 'DB schema (DDL), REST API endpoints, OpenAPI 3.1 spec, auth/RBAC, tenant isolation'),
    ('3', 'Frontend Developer', 'Continuous', 'Working web app consuming the design system + the live API contract'),
]
for i, row in enumerate(seq_data, start=1):
    for j, val in enumerate(row):
        c = seq.rows[i].cells[j]
        set_cell_margins(c)
        if i % 2 == 0:
            set_cell_bg(c, 'F7F7F7')
        c.text = ''
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(10)
        if j == 0:
            r.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ----- 3. Phase 1 — UI/UX Designer (First) -----
add_heading(doc, '3. Phase 1 — UI/UX Designer (First)', 1)

add_heading(doc, '3.1 Rationale (grounded in the SRS)', 2)
add_bullet(doc,
    "Part 11.3 of the SRS explicitly lists “Final visual design / brand kit for PDF templates and "
    "public website” as an unresolved OPEN ITEM. This must be closed before code is written; "
    "otherwise the frontend will produce inconsistent UI and the PDF module (Part 2.6.5) cannot "
    "render branded receipts, mark sheets, and certificates.")
add_bullet(doc,
    "Part 5 (UI Requirements) defines specifications only — screen names, navigation rules, and "
    "state behaviors. The actual mockups, component library, spacing, color tokens, and icon set "
    "do not yet exist and are a prerequisite for the frontend developer.")
add_bullet(doc,
    "The client constraint in §10.8 — “the system must not be extremely complicated, not difficult "
    "for teachers or accountants, and not dependent on one employee” — is a UX-driven constraint. "
    "These decisions must be made visually before any code is committed.")
add_bullet(doc,
    "The system has 40+ modules across Parts 2.1–2.7. Without a unified design system produced "
    "first, each frontend screen would diverge in density, navigation, and terminology.")
add_bullet(doc,
    "Multi-language support (Bangla, English, Arabic per §10.6) and Arabic RTL rendering on PDF "
    "require typography and layout decisions that only the designer can finalize.")
add_bullet(doc,
    "Mobile-first workflows — attendance capture in <60 seconds (§2.3.2), guardian portal "
    "(§2.2.3), marks entry (§5.4) — demand interactive prototypes to validate the interaction "
    "model before backend builds the supporting endpoints.")

add_heading(doc, '3.2 Deliverables', 2)
add_bullet(doc, "Brand kit: logo, color palette, typography stack for Bangla/English/Arabic.")
add_bullet(doc, "Design system: tokens (spacing, color, type scale), component library (forms, tables, modals, toasts, cards).")
add_bullet(doc, "Wireframes + hi-fi mockups for the Phase 0 + Phase 1 screens listed in Part 2 (Organization, RBAC, Audit, Security, Backup, Students, Attendance, Fees, Accounting, Dashboard).")
add_bullet(doc, "PDF/print templates for receipts, fee plans, mark sheets, result sheets, audit explorer exports.")
add_bullet(doc, "Empty / loading / error / permission-denied state designs (per §5.5).")
add_bullet(doc, "Mobile interaction prototypes for attendance and guardian portal.")

# ----- 4. Phase 2 — Backend Developer (Second) -----
add_heading(doc, '4. Phase 2 — Backend Developer (Second)', 1)

add_heading(doc, '4.1 Rationale (grounded in the SRS)', 2)
add_bullet(doc,
    "Part 6.7 states that the OpenAPI 3.1 specification is “generated from the code” and is the "
    "authoritative API contract. The frontend developer consumes this contract; therefore the "
    "backend must produce working endpoints before meaningful frontend integration can begin.")
add_bullet(doc,
    "Part 2.1 (Foundation Modules) is explicitly marked Phase 0 and must be implemented first "
    "because tenant isolation, branch scoping, RBAC, audit, and localization are cross-cutting "
    "concerns invoked by every subsequent module. These are backend responsibilities.")
add_bullet(doc,
    "Part 7 (Database Requirements) defines naming, multi-tenant isolation strategy (row-level "
    "security), indexing, and audit structures — but the actual DDL migrations must be produced "
    "by the backend developer before any data can flow.")
add_bullet(doc,
    "Part 3.6 (the “Golden Flow” financial posting contract) is a binding backend rule. Fees, "
    "donations, Zakat, salaries, and purchases all post through the Accounting ledger as balanced "
    "entries; the backend must enforce this or every finance module will diverge.")
add_bullet(doc,
    "Part 6.2 (Auth & Authorization): access/refresh token rotation, MFA, failed-login lockout, "
    "permission-code middleware, and row-level security in the repository layer are all backend "
    "infrastructure that the frontend depends on for any protected screen.")
add_bullet(doc,
    "Approximately 180+ API endpoints across all modules (Part 6.7) must exist before the "
    "frontend can do more than render static mockups.")

add_heading(doc, '4.2 Deliverables', 2)
add_bullet(doc, "Versioned, reversible DB migrations (DDL) implementing Part 7 rules — snake_case tables, uuid PKs, organization_id + branch_id columns, soft-delete, decimal money.")
add_bullet(doc, "Phase 0 module implementations: Organization & Multi-Branch, Module Configuration, Users & Roles (RBAC), Audit Trail, Security, Backup & Recovery.")
add_bullet(doc, "REST API following Part 6 conventions: /api/v1 prefix, standard list envelope, idempotency-key, error JSON, rate limiting.")
add_bullet(doc, "OpenAPI 3.1 spec auto-generated from code — this is the handoff artifact to the frontend.")
add_bullet(doc, "Auth: JWT access (15 min) + refresh (7 days, rotating), MFA (TOTP), permission-code middleware.")
add_bullet(doc, "Cross-cutting enforcement: tenant_id injection in repository layer, branch scoping, audit-on-write for §3.4 entities, fund isolation for Zakat (§3.7).")

# ----- 5. Phase 3 — Frontend Developer (Third) -----
add_heading(doc, '5. Phase 3 — Frontend Developer (Third)', 1)

add_heading(doc, '5.1 Rationale (grounded in the SRS)', 2)
add_bullet(doc,
    "The frontend’s job is to consume two artifacts: the design system (from the designer) and "
    "the live API contract / OpenAPI spec (from the backend). Building before either exists "
    "produces throwaway mock data and divergent UI.")
add_bullet(doc,
    "Part 5.1 requires permission-based UI: “menus, buttons, and fields the user cannot use are "
    "hidden; server still enforces authorization.” The frontend needs the permission-code catalog "
    "(e.g., fees.payment.create, students.notes.view) from the backend before it can wire "
    "visibility rules correctly.")
add_bullet(doc,
    "Part 6.5 (Idempotency-Key, async export jobs polled via /jobs/{id}) and Part 6.3 (standard "
    "list envelope) are integration patterns the frontend must implement against real endpoints — "
    "not approximations.")
add_bullet(doc,
    "Part 2.6.5 (Printing & PDF) and Part 5.6 require the frontend to trigger branded PDFs that "
    "render Bangla/English/Arabic correctly. The PDF templates come from the designer; the PDF "
    "engine endpoints come from the backend. The frontend cannot finalize this surface until both "
    "are ready.")
add_bullet(doc,
    "Mobile-first surfaces (§2.3.2 attendance, §2.2.3 guardian portal, §5.4) require both the "
    "mobile interaction prototypes (designer) and the underlying endpoints (backend) to be "
    "available before integration testing is meaningful.")

add_heading(doc, '5.2 Deliverables', 2)
add_bullet(doc, "Responsive web app shell: top bar (org/branch switcher, academic-year switcher, language switcher, notifications, user menu) and dynamic left navigation driven by enabled modules + permissions (§5.1).")
add_bullet(doc, "Component library implementation matching the designer’s tokens.")
add_bullet(doc, "Typed API client generated from the OpenAPI 3.1 spec handed over by the backend.")
add_bullet(doc, "Permission-aware UI: hide/disable based on permission codes; never trust client-only enforcement.")
add_bullet(doc, "Empty / loading (skeletons) / error / permission-denied states per §5.5.")
add_bullet(doc, "PDF trigger buttons and signed-URL download flow per §3.3 and §2.6.5.")

# ----- 6. Parallelization notes -----
add_heading(doc, '6. Safe Parallelization', 1)
add_para(doc,
    "Although the recommended first-pass order is strict (Designer → Backend → Frontend), two "
    "limited overlaps are safe once the designer has delivered the design system (week 2 onward):")
add_bullet(doc,
    "Backend Phase 0 (foundation) may begin while the designer is finishing hi-fi mockups for the "
    "Phase 1 screens — the foundation endpoints (auth, RBAC, organization) do not depend on "
    "screen-level design decisions.")
add_bullet(doc,
    "Frontend may begin scaffold work (shell, routing, design-system component wiring) once the "
    "design system is delivered, but should not implement module screens until the matching "
    "OpenAPI endpoints are merged on the backend.")
add_para(doc,
    "Any other parallelization risks rework and is discouraged for the first build pass.",
    italic=True)

# ----- 7. Handoff artifacts table -----
add_heading(doc, '7. Handoff Artifacts', 1)
add_para(doc,
    "The table below defines the binding artifact each role must produce and the role that "
    "consumes it. A role is not considered “done” until its handoff artifact is reviewed and "
    "accepted by the receiving role.")

ha = doc.add_table(rows=4, cols=3)
set_table_borders(ha)
hdr2 = ['From Role', 'Handoff Artifact', 'Consumed By']
for i, h in enumerate(hdr2):
    c = ha.rows[0].cells[i]
    set_cell_margins(c)
    set_cell_bg(c, '1F3A5F')
    c.text = ''
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

ha_data = [
    ('UI/UX Designer', 'Design system + mockups + PDF templates + brand kit', 'Frontend Developer (and Backend for PDF engine shape)'),
    ('Backend Developer', 'OpenAPI 3.1 spec + auth/RBAC + Phase 0 endpoints + DB schema', 'Frontend Developer'),
    ('Frontend Developer', 'Working web app integrated against live API + design system', 'QA / UAT (Parts 8 & 9)'),
]
for i, row in enumerate(ha_data, start=1):
    for j, val in enumerate(row):
        c = ha.rows[i].cells[j]
        set_cell_margins(c)
        if i % 2 == 0:
            set_cell_bg(c, 'F7F7F7')
        c.text = ''
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(10)
        if j == 0:
            r.bold = True

# ----- 8. Conclusion / Next steps -----
add_heading(doc, '8. Next Steps', 1)
add_para(doc,
    "Hand the SRS to the UI/UX Designer first with a brief covering the open visual design item in "
    "Part 11.3 and the screen catalog in Part 2 + Part 5. On delivery of the design system, hand "
    "the SRS to the Backend Developer with a brief covering Parts 2.1, 3, 6, and 7 — the goal is "
    "Phase 0 endpoints plus the OpenAPI 3.1 spec. Finally, hand the SRS to the Frontend Developer "
    "with the design system and the OpenAPI spec as binding inputs, covering Parts 2, 5, and 6. "
    "This sequence resolves the SRS open items in dependency order, produces a buildable API "
    "contract before frontend integration, and ensures the client’s “not complicated, not "
    "dependent on one employee” constraint is enforced from the first screen forward.")

doc.save(OUTPUT)
print(f"Saved: {OUTPUT}")
